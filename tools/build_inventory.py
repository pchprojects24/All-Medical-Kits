#!/usr/bin/env python3
"""Rebuild every derived inventory file from the SharePoint export.

Source      : Data/source/<SharePoint export>.xlsx  (one flat sheet covering
              every kit in the SharePoint medical kit catalogue)
Generated   : Data/inventory.json
              Data/kit_items.csv
              Data/master_inventory.csv
              Data/kits/<slug>.csv

The export is the only input and is authoritative for everything it lists:
NSN, description, unit of measure, quantity and price.

Run after replacing the export:

    python3 tools/build_inventory.py

Requires openpyxl (pip install openpyxl).
"""

import csv
import json
import os
import re
import sys
import unicodedata
from collections import Counter, OrderedDict

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE_DIR = os.path.join(ROOT, "Data", "source")
DATA_DIR = os.path.join(ROOT, "Data")
KIT_DIR = os.path.join(DATA_DIR, "kits")

# SharePoint export columns this build reads, by the English half of their
# bilingual header ("Item NSN / NNO de l'article" -> "item nsn"). Every other
# column in the export is ignored.
SHAREPOINT_COLUMNS = {
    "kit nsn": "kit_nsn",
    "kit name": "kit_name",
    "item nsn": "item_nsn",
    "item name": "description",
    "total qty": "total_qty",
    "uom": "uom",
    "total value": "total_value",
    "item type": "item_type",
    "# sub-kits req": "sub_kits_req",
    "sub-kit nsn": "sub_kit_nsn",
    "sub-kit name": "sub_kit_name",
    "full sub-kit path": "path",
    "controlled?": "controlled",
    "sg?": "sg",
    "storage": "storage",
    "on soa?": "soa",
    "o&pen keywords": "keywords",
}

NSN_RE = re.compile(r"^\d{4}-[0-9A-Z]{2}-[0-9A-Z]{3}-[0-9A-Z]{4}$")


def slugify(value):
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    value = re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-").lower()
    return re.sub(r"-{2,}", "-", value)


def short_slug(value, limit=50):
    """A slug cut at a word boundary, for kits named only by a long description."""
    slug = slugify(value)
    if len(slug) <= limit:
        return slug
    return slug[:limit + 1].rsplit("-", 1)[0]


def clean(value):
    """Normalise a cell to a trimmed string ('' for blanks)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def parse_number(value):
    """Parse a quantity/price cell. Returns None when not numeric."""
    text = clean(value).replace(",", "").replace("$", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def header_key(cell):
    """'Item NSN / NNO de l'article' -> 'item nsn'."""
    return re.split(r"\s+[/|]\s+", clean(cell), maxsplit=1)[0].strip().lower()


def clean_keywords(value):
    """' A,  B,  A ' -> 'A, B': the O&PEN keyword list, de-duplicated."""
    seen = OrderedDict()
    for word in clean(value).split(","):
        word = word.strip()
        if word:
            seen.setdefault(word.upper(), word)
    return ", ".join(seen.values())


# ── Read the export ────────────────────────────────────────────────────────

def read_sharepoint_rows():
    files = sorted(f for f in os.listdir(SOURCE_DIR)
                   if f.endswith(".xlsx") and not f.startswith("~$"))
    if len(files) != 1:
        sys.exit("expected exactly one SharePoint export in %s, found %s"
                 % (SOURCE_DIR, files))
    filename = files[0]
    book = openpyxl.load_workbook(os.path.join(SOURCE_DIR, filename),
                                  read_only=True, data_only=True)
    if len(book.worksheets) > 1:
        sys.exit("%s: expected one sheet, found %s" % (filename, book.sheetnames))
    sheet = book.worksheets[0]

    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    index = {}
    for position, cell in enumerate(header):
        key = header_key(cell)
        if key in SHAREPOINT_COLUMNS and key not in index:
            index[key] = position
    missing = sorted(set(SHAREPOINT_COLUMNS) - set(index))
    if missing:
        sys.exit("%s: missing expected columns: %s" % (filename, missing))

    rows = []
    for excel_row, raw in enumerate(sheet.iter_rows(min_row=2, values_only=True),
                                    start=2):
        row = {field: raw[index[key]] if index[key] < len(raw) else None
               for key, field in SHAREPOINT_COLUMNS.items()}
        # Blank spacer rows and the "Applied filters" footer have no item.
        if not clean(row["item_nsn"]):
            continue
        for field in ("kit_nsn", "kit_name", "item_nsn", "description", "uom",
                      "item_type", "sub_kit_nsn", "sub_kit_name", "path",
                      "controlled", "sg", "storage", "soa"):
            row[field] = clean(row[field])
        for field in ("kit_nsn", "item_nsn", "sub_kit_nsn", "uom"):
            row[field] = row[field].upper()
        row["keywords"] = clean_keywords(row["keywords"])
        row["total_qty"] = parse_number(row["total_qty"])
        row["total_value"] = parse_number(row["total_value"])
        row["sub_kits_req"] = parse_number(row["sub_kits_req"])
        if not row["kit_nsn"]:
            sys.exit("%s row %d: no Kit NSN" % (filename, excel_row))
        if row["total_qty"] is None:
            sys.exit("%s row %d: no Total Qty for %s" % (filename, excel_row,
                                                         row["item_nsn"]))
        row["source_row"] = excel_row
        rows.append(row)
    if not rows:
        sys.exit("%s: no item rows found" % filename)
    return filename, rows


def path_chain(row):
    """The kits from the row's kit down to its sub-kit.

    'A ➡️ B | NAME ➡️ C' -> ['A', 'B', 'C']. The part before '|' is the chain
    of NSNs above the sub-kit's parent; the sub-kit itself comes last.
    """
    head = row["path"].split("|", 1)[0]
    chain = [part.strip().upper() for part in head.split("➡️")
             if NSN_RE.match(part.strip().upper())]
    return chain + [row["sub_kit_nsn"]]


# ── Build ──────────────────────────────────────────────────────────────────

def build(sp_file, sp_rows):
    warnings = []

    # 1. Kits. Every kit in the catalogue is a "Kit NSN" with rows of its own;
    #    the export lists a kit's contents twice over: once as the kit itself
    #    (Sub-Kit blank, its direct contents) and again from every kit above it
    #    (Sub-Kit filled in). Only the direct rows become lines here; the rest
    #    give the hierarchy.
    kit_name, kit_order = {}, []
    for r in sp_rows:
        if r["kit_nsn"] not in kit_name:
            kit_name[r["kit_nsn"]] = r["kit_name"] or r["kit_nsn"]
            kit_order.append(r["kit_nsn"])

    # 2. Hierarchy. A kit can sit inside more than one parent (the same medic
    #    bag goes into several sets), so this is a graph, not a tree.
    #    copies[(K, …, C)] = how many C one K holds along that path, from
    #    "# Sub-Kits Req" on K's rows.
    parents = OrderedDict((nsn, []) for nsn in kit_order)
    copies = {}
    for r in sp_rows:
        if not r["sub_kit_nsn"]:
            continue
        chain = tuple(path_chain(r))
        if chain[0] != r["kit_nsn"]:
            sys.exit("%s row %d: path does not start at %s: %s"
                     % (sp_file, r["source_row"], r["kit_nsn"], r["path"]))
        for nsn in chain:
            if nsn not in kit_name:
                sys.exit("%s row %d: sub-kit %s has no rows of its own"
                         % (sp_file, r["source_row"], nsn))
        for parent, child in zip(chain, chain[1:]):
            if parent not in parents[child]:
                parents[child].append(parent)
        if r["sub_kits_req"] is not None:
            copies.setdefault(chain, Counter())[r["sub_kits_req"]] += 1

    def copies_along(chain):
        counts = copies.get(chain)
        if not counts:
            return None
        if len(counts) > 1:
            warnings.append("%s holds %s with inconsistent '# Sub-Kits Req' %s; "
                            "used the most common"
                            % (chain[0], " > ".join(chain[1:]), dict(counts)))
            copies[chain] = Counter({counts.most_common(1)[0][0]: 1})
        return counts.most_common(1)[0][0]

    def edge_qty(parent, child):
        """Copies of child in one parent."""
        direct = copies_along((parent, child))
        if direct is not None:
            return direct
        # The child appears only as a step in parent's paths (it holds sub-kits
        # but no items of its own): compare a grandchild's count both ways.
        for chain in list(copies):
            if len(chain) >= 3 and chain[:2] == (parent, child):
                in_parent = copies_along(chain)
                in_child = copies_along(chain[1:])
                if in_parent and in_child:
                    return in_parent / in_child
        warnings.append("no count for %s inside %s; assumed 1" % (child, parent))
        return 1.0

    # 3. Lines: a kit's own direct rows, plus one container line per sub-kit.
    lines = OrderedDict()
    merged = []
    for r in sp_rows:
        if r["sub_kit_nsn"]:
            continue
        key = (r["kit_nsn"], r["item_nsn"])
        if key in lines:
            # The same item listed twice in one kit (e.g. one case for the kit,
            # another for its resupply): one line, quantities added.
            line = lines[key]
            line["qty"] += r["total_qty"]
            merged.append(key)
            continue
        lines[key] = {
            "kit_nsn": r["kit_nsn"], "item_nsn": r["item_nsn"], "qty": r["total_qty"],
            "uom": r["uom"], "description": r["description"],
            "source_row": r["source_row"],
        }
    for child in kit_order:
        for parent in parents[child]:
            key = (parent, child)
            if key in lines:
                continue
            lines[key] = {
                "kit_nsn": parent, "item_nsn": child,
                "qty": edge_qty(parent, child), "uom": "KT",
                "description": kit_name[child], "source_row": None,
            }

    # 4. Item attributes and prices. SharePoint gives a total value per row, so
    #    the unit price is value / quantity, pooled across every row so rounding
    #    cannot make one item look differently priced in two kits.
    attrs, value_sum, qty_sum = {}, Counter(), Counter()
    for r in sp_rows:
        nsn = r["item_nsn"]
        a = attrs.setdefault(nsn, {"item_type": "", "controlled": False, "sg": False,
                                   "storage": "", "soa": False, "keywords": ""})
        a["item_type"] = a["item_type"] or r["item_type"]
        a["storage"] = a["storage"] or r["storage"]
        a["keywords"] = a["keywords"] or r["keywords"]
        a["controlled"] = a["controlled"] or r["controlled"].lower() == "yes"
        a["sg"] = a["sg"] or r["sg"].lower() == "yes"
        a["soa"] = a["soa"] or r["soa"].upper() == "SOA"
        value_sum[nsn] += r["total_value"] or 0
        qty_sum[nsn] += r["total_qty"]
    unit_price = {nsn: round(value_sum[nsn] / qty_sum[nsn], 2) if qty_sum[nsn] else None
                  for nsn in value_sum}
    for line in lines.values():
        line["price"] = unit_price.get(line["item_nsn"])

    # A container's price is the value of its contents, so price kits bottom-up.
    by_kit = {}
    for line in lines.values():
        by_kit.setdefault(line["kit_nsn"], []).append(line)
    kit_value = {}

    def value_of(nsn, seen=()):
        if nsn in kit_value:
            return kit_value[nsn]
        if nsn in seen:
            sys.exit("kit hierarchy loops through %s" % nsn)
        total = 0.0
        for line in by_kit.get(nsn, []):
            if line["item_nsn"] in kit_name:
                line["price"] = round(value_of(line["item_nsn"], seen + (nsn,)), 2)
            if line["price"] is not None and line["qty"] is not None:
                total += line["qty"] * line["price"]
        kit_value[nsn] = round(total, 2)
        return kit_value[nsn]

    for nsn in kit_order:
        value_of(nsn)

    for line in lines.values():
        line["is_nsn"] = bool(NSN_RE.match(line["item_nsn"]))
        line["extended_price"] = (round(line["qty"] * line["price"], 2)
                                  if line["qty"] is not None
                                  and line["price"] is not None else None)

    # 5. Kit records.
    kits = OrderedDict()
    used_slugs = {}
    for nsn in kit_order:
        slug = short_slug(kit_name[nsn]) or nsn.lower()
        if used_slugs.get(slug, nsn) != nsn:
            slug = "%s-%s" % (slug, nsn.lower())
        used_slugs[slug] = nsn
        own = by_kit.get(nsn, [])
        kits[nsn] = {
            "nsn": nsn,
            "name": kit_name[nsn],
            "slug": slug,
            "parent_nsns": list(parents[nsn]),
            "line_count": len(own),
            "total_value": kit_value[nsn],
            "sub_kits": [l["item_nsn"] for l in own if l["item_nsn"] in kit_name],
        }

    def depth_of(nsn, seen=()):
        ups = [p for p in kits[nsn]["parent_nsns"] if p not in seen]
        if not ups:
            return 0
        return 1 + max(depth_of(p, seen + (nsn,)) for p in ups)

    for nsn in kits:
        kits[nsn]["depth"] = depth_of(nsn)

    if merged:
        warnings.append("%d items listed twice in the same kit; quantities added: %s"
                        % (len(merged), ", ".join("%s in %s" % (i, k)
                                                  for k, i in merged)))
    return kits, list(lines.values()), attrs, warnings


def build_items(rows, kits, attrs):
    """Collapse the line rows into one record per distinct item."""
    items = OrderedDict()
    for row in rows:
        item = items.get(row["item_nsn"])
        if item is None:
            a = attrs.get(row["item_nsn"], {})
            item = items[row["item_nsn"]] = {
                "nsn": row["item_nsn"],
                "description": row["description"],
                "uom": row["uom"],
                "item_type": a.get("item_type") or ("Kit" if row["item_nsn"] in kits else ""),
                "storage": a.get("storage", ""),
                "controlled": a.get("controlled", False),
                "sg": a.get("sg", False),
                "soa": a.get("soa", False),
                "keywords": a.get("keywords", ""),
                "is_nsn": row["is_nsn"],
                "is_kit": row["item_nsn"] in kits,
                "prices": [],
                "memberships": [],
            }
        if row["price"] is not None:
            item["prices"].append(row["price"])
        item["memberships"].append({
            "kit_nsn": row["kit_nsn"],
            "qty": row["qty"],
            "price": row["price"],
            "extended_price": row["extended_price"],
        })

    for item in items.values():
        prices = item.pop("prices")
        item["price"] = round(prices[0], 2) if prices else None
        item["price_varies"] = len(set(prices)) > 1
        item["total_qty"] = round(sum(m["qty"] for m in item["memberships"]
                                      if m["qty"] is not None), 4)
        item["kit_count"] = len(item["memberships"])
    return items


def write_csv(path, header, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)


def fmt(value):
    """Format a number for CSV: integers without a decimal tail."""
    if value is None:
        return ""
    if isinstance(value, float):
        value = round(value, 4)
        if value.is_integer():
            return str(int(value))
    return str(value)


def yn(flag):
    return "Y" if flag else ""


def main():
    sp_file, sp_rows = read_sharepoint_rows()
    kits, rows, attrs, warnings = build(sp_file, sp_rows)
    items = build_items(rows, kits, attrs)

    ordered = sorted(rows, key=lambda r: (kits[r["kit_nsn"]]["name"], r["item_nsn"]))

    def line_cells(r):
        item = items[r["item_nsn"]]
        return [r["item_nsn"], r["description"], fmt(r["qty"]), r["uom"],
                item["item_type"], fmt(r["price"]), fmt(r["extended_price"]),
                yn(item["controlled"]), item["storage"], yn(r["item_nsn"] in kits)]

    line_header = ["Item NSN", "Description", "Quantity", "Unit of Measure", "Item Type",
                   "Unit Price", "Extended Price", "Controlled", "Storage", "Is Sub-Kit"]

    # 1. Every kit-item line, the canonical flat table.
    write_csv(
        os.path.join(DATA_DIR, "kit_items.csv"),
        ["Kit NSN", "Kit Name"] + line_header,
        [[r["kit_nsn"], kits[r["kit_nsn"]]["name"]] + line_cells(r) for r in ordered],
    )

    # 2. One file per kit. Remove any left over from a renamed or deleted kit
    #    first: an orphan is indistinguishable from current data.
    os.makedirs(KIT_DIR, exist_ok=True)
    expected = {kit["slug"] + ".csv" for kit in kits.values()}
    for stale in sorted(set(os.listdir(KIT_DIR)) - expected):
        if stale.endswith(".csv"):
            os.remove(os.path.join(KIT_DIR, stale))
            print("removed stale %s" % stale)
    by_kit = {}
    for r in ordered:
        by_kit.setdefault(r["kit_nsn"], []).append(r)
    for kit in kits.values():
        write_csv(os.path.join(KIT_DIR, kit["slug"] + ".csv"), line_header,
                  [line_cells(r) for r in by_kit.get(kit["nsn"], [])])

    # 3. Item-level rollup, one row per distinct item across all kits.
    #    "Is Kit" marks containers that have contents of their own: their value
    #    is already the sum of their contents, so anyone counting stock or money
    #    should filter them out.
    write_csv(
        os.path.join(DATA_DIR, "master_inventory.csv"),
        ["NSN", "Description", "Unit of Measure", "Total Quantity", "Kit Count",
         "Unit Price", "Item Type", "Controlled", "SG", "Storage", "On SOA", "Is Kit",
         "Kit Membership"],
        [[item["nsn"], item["description"], item["uom"], fmt(item["total_qty"]),
          item["kit_count"], fmt(item["price"]), item["item_type"],
          yn(item["controlled"]), yn(item["sg"]), item["storage"], yn(item["soa"]),
          yn(item["is_kit"]),
          "; ".join("%s (%s) x%s" % (kits[m["kit_nsn"]]["name"], m["kit_nsn"],
                                     fmt(m["qty"]))
                    for m in item["memberships"])]
         for item in sorted(items.values(), key=lambda i: i["nsn"])],
    )

    # 4. Structured data for the browser app. Compact (no indentation): at
    #    this size the whitespace alone would be a large share of the download.
    payload = {
        "generated_from": "Data/source/%s" % sp_file,
        "kit_count": len(kits),
        "item_count": len(items),
        "line_count": len(rows),
        "kits": list(kits.values()),
        "items": sorted(items.values(), key=lambda i: i["nsn"]),
    }
    with open(os.path.join(DATA_DIR, "inventory.json"), "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
        handle.write("\n")

    roots = sum(1 for k in kits.values() if not k["parent_nsns"])
    print("kits: %d (%d top-level)  items: %d  lines: %d"
          % (len(kits), roots, len(items), len(rows)))
    for warning in warnings:
        print("note: " + warning)

    # Keep index.html's asset versions in step, so browsers never pair a cached
    # old app.js with a newer page.
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import stamp_assets
    stamp_assets.main()


if __name__ == "__main__":
    main()
